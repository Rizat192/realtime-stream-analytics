import json
import os
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu, Inches, Pt

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from src.config import load_config, path_in

DARK = RGBColor(0x1F, 0x3B, 0x57)
GREY = RGBColor(0x55, 0x5F, 0x6B)
ACCENT = RGBColor(0x1F, 0x77, 0xB4)
WIDTH = Inches(13.333)
HEIGHT = Inches(7.5)


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def textbox(slide, left, top, width, height, text, size=18, bold=False, color=GREY):
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.word_wrap = True
    lines = text if isinstance(text, list) else [text]
    for index, line in enumerate(lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = line
        paragraph.space_after = Pt(8)
        for run in paragraph.runs:
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = color
            run.font.name = "Calibri"
    return box


def heading(slide, text):
    textbox(slide, Inches(0.6), Inches(0.35), Inches(12.1), Inches(0.9),
            text, size=28, bold=True, color=DARK)
    line = slide.shapes.add_shape(1, Inches(0.6), Inches(1.15), Inches(12.1), Emu(20000))
    line.fill.solid()
    line.fill.fore_color.rgb = ACCENT
    line.line.fill.background()
    line.shadow.inherit = False


def picture(slide, path, left, top, width):
    if os.path.exists(path):
        slide.shapes.add_picture(path, left, top, width=width)


def bullet_slide(prs, title, lines, size=18):
    slide = blank(prs)
    heading(slide, title)
    textbox(slide, Inches(0.7), Inches(1.5), Inches(11.9), Inches(5.4), lines, size=size)
    return slide


def figure_slide(prs, title, image, lines=None, image_width=Inches(8.6)):
    slide = blank(prs)
    heading(slide, title)
    if lines:
        textbox(slide, Inches(0.7), Inches(1.45), Inches(3.4), Inches(5.3), lines, size=15)
        picture(slide, image, Inches(4.35), Inches(1.5), image_width)
    else:
        picture(slide, image, Inches(1.9), Inches(1.6), Inches(9.6))
    return slide


def table_slide(prs, title, headers, rows, note=None, width=Inches(12.0)):
    slide = blank(prs)
    heading(slide, title)
    shape = slide.shapes.add_table(len(rows) + 1, len(headers), Inches(0.65), Inches(1.5),
                                   width, Inches(0.4 + 0.36 * len(rows)))
    table = shape.table
    for column, text in enumerate(headers):
        cell = table.cell(0, column)
        cell.text = str(text)
        for paragraph in cell.text_frame.paragraphs:
            for run in paragraph.runs:
                run.font.size = Pt(14)
                run.font.bold = True
    for index, row in enumerate(rows, start=1):
        for column, value in enumerate(row):
            cell = table.cell(index, column)
            cell.text = str(value)
            for paragraph in cell.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(13)
    if note:
        textbox(slide, Inches(0.7), Inches(6.55), Inches(11.9), Inches(0.7), note, size=14)
    return slide


def fmt(value, digits=3):
    return format(round(float(value), digits), "." + str(digits) + "f")


def main():
    cfg = load_config()
    with open(path_in(cfg, "results", "offline_results.json"), "r", encoding="utf-8") as handle:
        offline = json.load(handle)
    with open(path_in(cfg, "results", "stream_summary.json"), "r", encoding="utf-8") as handle:
        stream = json.load(handle)

    prs = Presentation()
    prs.slide_width = WIDTH
    prs.slide_height = HEIGHT

    slide = blank(prs)
    textbox(slide, Inches(0.9), Inches(2.3), Inches(11.5), Inches(1.6),
            "Real-Time Monitoring, Time-Series Analysis and Stream Intelligence",
            size=38, bold=True, color=DARK)
    textbox(slide, Inches(0.9), Inches(3.9), Inches(11.5), Inches(1.0),
            "Endterm project: API traffic stream, anomaly detection, real-time ML, "
            "distributed processing", size=20)
    textbox(slide, Inches(0.9), Inches(5.0), Inches(11.5), Inches(1.2),
            ["{:,} events, {:,} one-minute windows, {} labelled incidents".format(
                int(stream["events_processed"]), offline["dataset"]["minute_windows"],
                cfg["generator"]["anomalies"]["count"]),
             "{} concurrent processes connected over TCP, no shared memory".format(
                 stream["processes"])], size=17)

    bullet_slide(prs, "Problem and objectives", [
        "A stream of API request records carries timestamp, endpoint, region, latency, status and size.",
        "Objective 1. Continuously updated visualization with more than one view.",
        "Objective 2. Time-series analysis with rolling windows, resampling, trend and seasonality.",
        "Objective 3. Anomaly detection that is implemented and evaluated against ground truth.",
        "Objective 4. Real-time machine learning using only information available at prediction time.",
        "Objective 5. Distributed stream processing with measured operational characteristics.",
    ])

    stats = offline["dataset"]
    table_slide(prs, "The data stream", ["Property", "Value"], [
        ["Events", "{:,}".format(int(stream["events_processed"]))],
        ["Windows", "{:,} one-minute tumbling windows over 3 days".format(stats["minute_windows"])],
        ["Endpoints", "6, each with its own rate, latency and error profile"],
        ["Mean traffic", "{:.1f} requests per minute".format(stats["mean_requests_per_minute"])],
        ["Injected incidents", "{} incidents, {:,} labelled minutes ({:.1%})".format(
            cfg["generator"]["anomalies"]["count"], stats["anomalous_minutes"], stats["anomaly_rate"])],
        ["Incident types", "latency spike, error burst, traffic drop, traffic surge"],
    ], note="The generator is seeded, so the dataset, the models and every number are reproducible.")

    figure_slide(prs, "Architecture", path_in(cfg, "figures", "fig_architecture.png"))

    figure_slide(prs, "Live dashboard, four views", path_in(cfg, "figures", "fig_dashboard.png"), [
        "Volume against its 15 minute rolling mean.",
        "Tail latency with alarms as they fire.",
        "Error rate against predicted incident risk.",
        "Throughput and per-stage backlog.",
        "Redrawn by the sink process as scored windows arrive.",
    ])

    summary = offline["timeseries"]
    figure_slide(prs, "Time-series analysis", path_in(cfg, "figures", "fig_stl_decomposition.png"), [
        "One-minute resampling of an irregular event stream.",
        "Rolling windows of 15 and 60 minutes.",
        "STL with a daily period of 1440 minutes.",
        "Seasonal strength {:.2f}, trend strength {:.2f}.".format(
            summary["seasonal_strength"], summary["trend_strength"]),
        "Autocorrelation at one day {:.2f}.".format(summary["acf_lag_1440"]),
        "Busiest hour {}:00, quietest hour {}:00.".format(summary["peak_hour"], summary["trough_hour"]),
    ])

    rows = [[item["method"], fmt(item["precision"], 2), fmt(item["recall"], 2), fmt(item["f1"], 2),
             fmt(item["roc_auc"], 2), "{}/{}".format(item["events_detected"], item["anomaly_events"])]
            for item in offline["anomaly"]]
    detector = stream["streaming_detector"]
    rows.append(["streaming robust z, live run", fmt(detector["precision"], 2),
                 fmt(detector["recall"], 2), fmt(detector["f1"], 2), fmt(detector["roc_auc"], 2),
                 "{}/{}".format(detector["events_detected"], detector["anomaly_events"])])
    table_slide(prs, "Anomaly detection, evaluated on a held-out window",
                ["Method", "Precision", "Recall", "F1", "ROC AUC", "Incidents caught"], rows,
                note="Robust z score reacts to every incident, isolation forest raises fewer "
                     "but cleaner alarms.")

    figure_slide(prs, "Detector behaviour", path_in(cfg, "figures", "fig_anomaly_detection.png"))

    rows = [[item["model"], fmt(item["precision"], 2), fmt(item["recall"], 2), fmt(item["f1"], 2),
             fmt(item["roc_auc"], 2), fmt(item["pr_auc"], 2)] for item in offline["classification"]]
    table_slide(prs, "Real-time classification: is the next minute an incident",
                ["Model", "Precision", "Recall", "F1", "ROC AUC", "PR AUC"], rows,
                note="Time-ordered split, {:,} train and {:,} test windows, {} features, all "
                     "computed from closed windows.".format(
                         offline["ml_info"]["train_rows"], offline["ml_info"]["test_rows"],
                         offline["ml_info"]["features"]))

    onset_rows = [[item["model"], item["onset_minutes"], fmt(item["onset_recall"], 2),
                   fmt(item["ongoing_recall"], 2)] for item in offline.get("onset", [])]
    if onset_rows:
        table_slide(prs, "Where the skill actually comes from",
                    ["Model", "Onset minutes", "Recall at onset", "Recall while ongoing"], onset_rows,
                    note="Most of the F1 comes from incidents that are already visible. Predicting "
                         "the first minute of an incident remains the hard part.")

    table_slide(prs, "Distributed processing, measured", ["Operational characteristic", "Value"], [
        ["Processes", "{} across 6 stages".format(stream["processes"])],
        ["Partitioning", "{} queues keyed by CRC32 of the endpoint".format(cfg["pipeline"]["partitions"])],
        ["Throughput", "{:,.0f} events per second mean, {:,.0f} peak".format(
            stream["mean_throughput_eps"], stream["peak_throughput_eps"])],
        ["End-to-end latency", "{:.0f} ms mean, {:.0f} ms p95".format(
            stream["end_to_end_ms_mean"], stream["end_to_end_ms_p95"])],
        ["Events and windows", "{:,} events, {:,} windows completed".format(
            int(stream["events_processed"]), stream["simulated_minutes"])],
        ["Backlog", "{:.1f} mean, {:,} peak at the {} stage".format(
            stream["backlog_mean"], stream["backlog_peak"],
            stream.get("backlog_peak_stage", "scoring"))],
        ["Late records dropped", "{:,} behind a {} minute watermark".format(
            stream.get("late_events_dropped", 0), stream.get("watermark_minutes", 20))],
        ["Replay", "{:.1f} hours of traffic in {:.0f} seconds".format(
            stream["simulated_hours"], stream["wall_clock_seconds"])],
    ])

    figure_slide(prs, "Operational metrics during the run",
                 path_in(cfg, "figures", "fig_ops_metrics.png"))

    bullet_slide(prs, "Engineering: finding the bottleneck", [
        "Per-stage traces showed the ML worker at about 30 ms per window, almost all of it the pandas feature build.",
        "Its inbox grew without bound while every other stage sat idle.",
        "The same feature row is now computed with NumPy on the trailing buffers, about 50 times faster.",
        "Producers also drifted, so late batches reopened closed windows and minutes were scored twice. Closed windows now stay closed.",
    ])

    bullet_slide(prs, "Limitations, stated honestly", [
        "The data is synthetic, so incident shapes are cleaner than production traffic.",
        "The distributed layer runs on one machine: no broker failures, no rebalancing, no network faults.",
        "Models are trained once and served: no drift detection, no scheduled retraining.",
        "Supervised labels come from the generator and would not exist at prediction time in production.",
        "The held-out window contains {} incidents, which is a small sample for event-level claims.".format(
            offline["anomaly"][0]["anomaly_events"]),
    ])

    bullet_slide(prs, "Conclusion and demonstration", [
        "All five required capabilities are implemented, measured and reproducible from one command.",
        "pip install -r requirements.txt, then python run_all.py.",
        "python scripts/run_stream.py replays {:.1f} hours of traffic in about {:.0f} seconds with "
        "the live dashboard.".format(stream["simulated_hours"], stream["wall_clock_seconds"]),
        "Metrics are written to results/, figures to figures/, the report to report/.",
    ])

    out = os.path.join(cfg["root"], "presentation", "Endterm_Presentation.pptx")
    prs.save(out)
    print("saved " + out)


if __name__ == "__main__":
    main()
